"""Reproducible 1000-page acceptance drill for Vedomo.

The default ``extract`` mode is cheap: it builds a temporary DOCX, runs the
real document loader and proves that control facts from the first through the
last page survive extraction. ``retrieval`` adds real embedding/reranking;
``full`` also calls the configured LLM.

Examples:
    python -m scripts.large_document_benchmark
    python -m scripts.large_document_benchmark --mode retrieval
    python -m scripts.large_document_benchmark --mode full --repeat 3
"""

from __future__ import annotations

import argparse
import json
import sys
import tempfile
import time
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


CONTROL_FACTS = {
    1: "Код запуска линии Север — NS-1041.",
    250: "Перед заменой фильтра Q250 питание отключают не менее чем на 17 минут.",
    500: "Предельная рабочая температура модуля M500 составляет 63 градуса Цельсия.",
    750: "Контрольная последовательность узла K750: синий, янтарный, белый.",
    1000: "Гарантийный срок блока Z1000 составляет 47 месяцев.",
}


def _fact_for_page(page: int, pages: int) -> str | None:
    if page in CONTROL_FACTS and page <= pages:
        return CONTROL_FACTS[page]
    return None


def build_document(path: Path, pages: int) -> None:
    document = Document()
    base = (
        "Этот раздел описывает штатную эксплуатацию учебного оборудования, "
        "проверку состояния, порядок подготовки рабочего места и регистрацию "
        "результатов. Формулировка намеренно нейтральна и служит нагрузочным "
        "текстом для проверки полного прохождения документа через загрузчик."
    )
    for page in range(1, pages + 1):
        document.add_heading(f"РАЗДЕЛ {page:04d}", level=1)
        fact = _fact_for_page(page, pages)
        if fact:
            document.add_paragraph(fact)
        for paragraph in range(1, 8):
            document.add_paragraph(
                f"Страница {page:04d}, абзац {paragraph}: {base} "
                f"Контрольная метка страницы — P{page:04d}-{paragraph}."
            )
        if page != pages:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.save(path)


def extraction_drill(path: Path, pages: int) -> dict:
    import config
    from src.document_loader import load_file

    started = time.perf_counter()
    text = load_file(str(path))
    elapsed = time.perf_counter() - started
    expected = {page: fact for page, fact in CONTROL_FACTS.items() if page <= pages}
    missing = [page for page, fact in expected.items() if fact not in text]
    return {
        "mode": "extract",
        "pages": pages,
        "file_bytes": path.stat().st_size,
        "extracted_chars": len(text),
        "max_upload_bytes": config.MAX_UPLOAD_BYTES,
        "max_index_chars": config.MAX_INDEX_CHARS,
        "within_upload_limit": path.stat().st_size <= config.MAX_UPLOAD_BYTES,
        "within_index_limit": len(text) <= config.MAX_INDEX_CHARS,
        "control_pages": sorted(expected),
        "missing_control_pages": missing,
        "seconds": round(elapsed, 3),
        "passed": not missing,
    }


def _dataset(path: Path):
    from evals.dataset import EvalCase, EvalDataset

    filename = path.name
    cases = [
        EvalCase(
            id="first-page-code",
            question="Какой код запуска у линии Север?",
            material=filename,
            expected_keywords=["NS-1041"],
            expected_context=["NS-1041"],
            category="first_page",
            critical=True,
        ),
        EvalCase(
            id="quarter-indirect-wait",
            question="Можно ли приступать к замене фильтра Q250 через десять минут после отключения?",
            material=filename,
            expected_keywords=["17"],
            expected_context=["не менее чем на 17 минут"],
            forbidden_keywords=["можно приступать через десять минут"],
            category="indirect_constraint",
            critical=True,
        ),
        EvalCase(
            id="middle-temperature",
            question="Назови предельную рабочую температуру M500.",
            material=filename,
            expected_keywords=["63"],
            expected_context=["63 градуса Цельсия"],
            category="middle_page",
            critical=True,
        ),
        EvalCase(
            id="three-quarter-sequence",
            question="В каком порядке идут цвета для контроля узла K750?",
            material=filename,
            expected_keywords=["синий", "янтарный", "белый"],
            expected_context=["синий, янтарный, белый"],
            category="late_page",
            critical=True,
        ),
        EvalCase(
            id="last-page-warranty",
            question="Сколько месяцев гарантии у блока Z1000?",
            material=filename,
            expected_keywords=["47"],
            expected_context=["47 месяцев"],
            category="last_page",
            critical=True,
        ),
        EvalCase(
            id="unknown-mass",
            question="Какова масса блока Z1000 в килограммах?",
            material=filename,
            expect_refusal=True,
            category="refusal",
            critical=True,
        ),
    ]
    available = []
    for case in cases:
        page = {"first-page-code": 1, "quarter-indirect-wait": 250,
                "middle-temperature": 500, "three-quarter-sequence": 750,
                "last-page-warranty": 1000}.get(case.id, 1)
        if page <= int(path.stem.rsplit("-", 1)[-1]):
            available.append(case)
    return EvalDataset(
        name=f"large-document-{path.stem.rsplit('-', 1)[-1]}-pages",
        description="Synthetic end-to-end large-document acceptance drill.",
        materials=[str(path)],
        cases=available,
    )


def rag_drill(path: Path, pages: int, *, mode: str, repeat: int) -> dict:
    from evals.runner import report_to_dict, run_dataset

    dataset = _dataset(path)
    report = run_dataset(
        dataset,
        repeat=repeat,
        retrieval_only=mode == "retrieval",
        stream=mode == "full",
    )
    result = report_to_dict(report)
    result["requested_pages"] = pages
    return result


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is not None:
            try:
                reconfigure(errors="replace")
            except (ValueError, OSError):
                pass
    parser = argparse.ArgumentParser(description="Run the Vedomo large-document drill.")
    parser.add_argument("--pages", type=int, default=1000)
    parser.add_argument("--mode", choices=("extract", "retrieval", "full"), default="extract")
    parser.add_argument("--repeat", type=int, default=1)
    args = parser.parse_args(argv)
    if args.pages < 1:
        parser.error("--pages must be positive")
    if args.repeat < 1:
        parser.error("--repeat must be positive")

    generated_started = time.perf_counter()
    with tempfile.TemporaryDirectory(prefix="vedomo-large-doc-") as temporary:
        path = Path(temporary) / f"vedomo-benchmark-{args.pages}.docx"
        build_document(path, args.pages)
        generation_seconds = time.perf_counter() - generated_started
        if args.mode == "extract":
            result = extraction_drill(path, args.pages)
        else:
            result = rag_drill(path, args.pages, mode=args.mode, repeat=args.repeat)
        result["generation_seconds"] = round(generation_seconds, 3)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        passed = result.get("passed")
        if passed is None:
            summary = result.get("summary") or {}
            passed = summary.get("passed") == summary.get("n")
        return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
