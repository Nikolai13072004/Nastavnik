"""Stage 18-3: assignments — create / publish / take / submit / results.

Drives the assignment lifecycle through the API as a real teacher (the
``authed_client`` owner) plus joined students. No KB/LLM: the teacher posts
curated questions directly (generation is the trainer's job, tested separately).

The security-critical guarantees are asserted here: a draft is invisible to
students, correct answers are withheld until submit, the score is computed
server-side, and a student can't reach the manage routes or another course's
assignment.
"""

import api_app  # noqa: F401  (registers models before fixtures' create_all)
from conftest import WorkspaceClient as TestClient

# Two questions with known answers: q0 → index 1, q1 → index 0.
_QS = [
    {"question": "2+2?", "options": ["3", "4", "5", "6"], "correct_index": 1, "explanation": "четыре"},
    {"question": "Столица РФ?", "options": ["Москва", "Питер"], "correct_index": 0, "explanation": ""},
]


def _register(client, email):
    return client.post(
        "/api/auth/register",
        json={"email": email, "password": "password12345", "display_name": email.split("@")[0]},
    )


def _owner_course_with_assignment(authed_client, publish=True):
    """Owner creates a course (auto-switched in) + a quiz assignment; returns
    (course_id, join_code, assignment_id)."""
    course_id = authed_client.post("/api/courses", json={"name": "Сети"}).json()["id"]
    code = authed_client.get(f"/api/courses/{course_id}").json()["join_code"]
    created = authed_client.post("/api/assignments", json={"title": "Контрольная 1", "questions": _QS})
    assert created.status_code == 200, created.text
    aid = created.json()["id"]
    if publish:
        pub = authed_client.post(f"/api/assignments/{aid}/publish", json={"published": True})
        assert pub.status_code == 200
    return course_id, code, aid


def test_create_requires_title_and_questions(authed_client):
    authed_client.post("/api/courses", json={"name": "Сети"})
    assert authed_client.post("/api/assignments", json={"title": "", "questions": _QS}).status_code == 400
    assert authed_client.post("/api/assignments", json={"title": "X", "questions": []}).status_code == 400


def test_draft_hidden_from_students_until_published(authed_client):
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=False)

    owner_list = authed_client.get("/api/assignments").json()
    assert owner_list["can_manage"] is True
    assert any(it["id"] == aid and it["is_published"] is False for it in owner_list["assignments"])

    with TestClient(api_app.app) as student:
        _register(student, "draft-stud@example.com")
        student.post("/api/courses/join", json={"code": code})

        slist = student.get("/api/assignments").json()
        assert slist["can_manage"] is False
        assert slist["assignments"] == []  # draft hidden
        assert student.get(f"/api/assignments/{aid}").status_code == 404

        # Owner publishes → it appears for the student.
        authed_client.post(f"/api/assignments/{aid}/publish", json={"published": True})
        slist2 = student.get("/api/assignments").json()
        assert any(it["id"] == aid for it in slist2["assignments"])


def test_answers_hidden_until_submit_and_scored_server_side(authed_client):
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=True)

    with TestClient(api_app.app) as student:
        _register(student, "score-stud@example.com")
        student.post("/api/courses/join", json={"code": code})

        view = student.get(f"/api/assignments/{aid}").json()
        assert view["submitted"] is False
        assert view["result"] is None
        assert len(view["questions"]) == 2
        for q in view["questions"]:  # no answer leaked before submitting
            assert "correct_index" not in q
            assert "explanation" not in q

        res = student.post(f"/api/assignments/{aid}/submit", json={"answers": [1, 0]}).json()
        assert res["score"] == 2 and res["total"] == 2
        assert res["questions"][0]["correct_index"] == 1  # revealed after submit

        after = student.get(f"/api/assignments/{aid}").json()
        assert after["submitted"] is True
        assert after["result"]["score"] == 2


def test_client_cannot_spoof_score(authed_client):
    """Even if the client posts garbage, the score comes from the server."""
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=True)
    with TestClient(api_app.app) as student:
        _register(student, "spoof-stud@example.com")
        student.post("/api/courses/join", json={"code": code})
        # One right (q0=1), one wrong (q1 should be 0).
        res = student.post(f"/api/assignments/{aid}/submit", json={"answers": [1, 1]}).json()
        assert res["score"] == 1


def test_resubmit_is_ignored_one_attempt(authed_client):
    """One attempt only (Stage 46): after the first submit, a resubmit must NOT
    re-score or overwrite - otherwise a student could submit blind, read the
    answers revealed on the result screen, then resubmit them for 100%."""
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=True)
    with TestClient(api_app.app) as student:
        _register(student, "re-stud@example.com")
        student.post("/api/courses/join", json={"code": code})
        assert student.post(f"/api/assignments/{aid}/submit", json={"answers": [0, 1]}).json()["score"] == 0
        # Resubmitting the now-revealed correct answers must NOT improve the score.
        assert student.post(f"/api/assignments/{aid}/submit", json={"answers": [1, 0]}).json()["score"] == 0

    results = authed_client.get(f"/api/assignments/{aid}/results").json()
    assert results["attempts_count"] == 1  # still a single attempt
    assert results["attempts"][0]["score"] == 0  # the first attempt's score stands


def test_student_cannot_manage(authed_client):
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=True)
    with TestClient(api_app.app) as student:
        _register(student, "noman-stud@example.com")
        student.post("/api/courses/join", json={"code": code})
        assert student.post("/api/assignments", json={"title": "X", "questions": _QS}).status_code == 403
        assert student.post(f"/api/assignments/{aid}/publish", json={"published": False}).status_code == 403
        assert student.delete(f"/api/assignments/{aid}").status_code == 403
        assert student.get(f"/api/assignments/{aid}/results").status_code == 403


def test_results_analytics(authed_client):
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=True)

    with TestClient(api_app.app) as s1:
        _register(s1, "a1@example.com")
        s1.post("/api/courses/join", json={"code": code})
        s1.post(f"/api/assignments/{aid}/submit", json={"answers": [1, 0]})  # 2/2 = 100%
    with TestClient(api_app.app) as s2:
        _register(s2, "a2@example.com")
        s2.post("/api/courses/join", json={"code": code})
        s2.post(f"/api/assignments/{aid}/submit", json={"answers": [1, 1]})  # 1/2 = 50%

    results = authed_client.get(f"/api/assignments/{aid}/results").json()
    assert results["students_total"] == 2
    assert results["attempts_count"] == 2
    assert results["avg_score_pct"] == 75.0  # mean(100, 50)

    qstats = {qs["index"]: qs for qs in results["question_stats"]}
    assert qstats[0]["accuracy_pct"] == 100.0  # both got q0
    assert qstats[1]["accuracy_pct"] == 50.0   # one got q1
    # The teacher's results view includes the full questions (with answers).
    assert len(results["questions"]) == 2
    assert results["questions"][0]["correct_index"] == 1


def test_outsider_cannot_reach_assignment(authed_client):
    _course_id, _code, aid = _owner_course_with_assignment(authed_client, publish=True)
    with TestClient(api_app.app) as outsider:
        _register(outsider, "out@example.com")  # never joins → active = personal
        assert outsider.get(f"/api/assignments/{aid}").status_code == 404
        assert outsider.post(f"/api/assignments/{aid}/submit", json={"answers": [1, 0]}).status_code == 404
        assert outsider.get("/api/assignments").json()["assignments"] == []


def test_delete_assignment_removes_it(authed_client):
    _course_id, _code, aid = _owner_course_with_assignment(authed_client, publish=True)
    assert authed_client.delete(f"/api/assignments/{aid}").status_code == 200
    assert all(it["id"] != aid for it in authed_client.get("/api/assignments").json()["assignments"])
    assert authed_client.get(f"/api/assignments/{aid}").status_code == 404


def _load_results_sheet(content):
    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content))
    return wb.active


def test_results_xlsx_export(authed_client):
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=True)
    with TestClient(api_app.app) as s1:
        _register(s1, "csv1@example.com")
        s1.post("/api/courses/join", json={"code": code})
        s1.post(f"/api/assignments/{aid}/submit", json={"answers": [1, 0]})  # 2/2 = 100%

    resp = authed_client.get(f"/api/assignments/{aid}/results.xlsx")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    assert "attachment" in resp.headers["content-disposition"]
    assert resp.headers["content-disposition"].rstrip('"').endswith(".xlsx")

    ws = _load_results_sheet(resp.content)
    header = [cell.value for cell in ws[1]]
    assert header == ["Студент", "Email", "Балл", "Всего", "Процент", "Сдано"]
    rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]
    assert any(
        r[1] == "csv1@example.com" and r[2] == 2 and r[3] == 2 and r[4] == "100%" for r in rows
    )


def test_results_xlsx_student_forbidden(authed_client):
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=True)
    with TestClient(api_app.app) as student:
        _register(student, "csv-stud@example.com")
        student.post("/api/courses/join", json={"code": code})
        assert student.get(f"/api/assignments/{aid}/results.xlsx").status_code == 403


def test_results_xlsx_neutralizes_formula_injection(authed_client):
    """A student controls their own display_name; a formula-leading value must be
    neutralized so it can't auto-run in the teacher's spreadsheet (formula injection)."""
    _course_id, code, aid = _owner_course_with_assignment(authed_client, publish=True)
    with TestClient(api_app.app) as student:
        student.post(
            "/api/auth/register",
            json={
                "email": "evil-csv@example.com",
                "password": "password12345",
                "display_name": "=1+2",
            },
        )
        student.post("/api/courses/join", json={"code": code})
        student.post(f"/api/assignments/{aid}/submit", json={"answers": [1, 0]})

    ws = _load_results_sheet(authed_client.get(f"/api/assignments/{aid}/results.xlsx").content)
    names = [row[0].value for row in ws.iter_rows(min_row=2)]
    # The malicious name is present, but prefixed so the cell renders as text.
    assert "'=1+2" in names
    # Defense in depth: no string cell may start with a raw formula leader.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("=", "+", "-", "@")), cell.value
