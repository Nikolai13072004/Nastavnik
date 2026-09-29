from pathlib import Path

import pytest

import config
from src import document_loader, ocr
from src.document_loader import (
    UnsupportedFormatError,
    detect_format,
    load_file,
    load_rtf,
    sniff_format,
)
from src.ocr import ExtractionUnavailableError


def test_load_txt_file():
    path = Path(__file__).parent / "sample.txt"
    text = load_file(str(path))

    assert "Vedomo" in text
    assert "тестовый учебный текст" in text


# --- Stage 38: .docx tables, .pptx, .rtf ----------------------------------

def test_supported_formats_include_new_ones():
    for ext in (".pptx", ".rtf", ".odt", ".ods", ".xlsx", ".csv", ".doc", ".png"):
        assert ext in config.SUPPORTED_FORMATS


def test_docx_now_extracts_table_text(tmp_path):
    from docx import Document

    path = tmp_path / "doc.docx"
    doc = Document()
    doc.add_paragraph("Обычный абзац про телеграф.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Год изобретения"
    table.rows[0].cells[1].text = "1837"
    doc.save(str(path))

    text = load_file(str(path))
    assert "Обычный абзац про телеграф" in text
    # The table cell was lost before Stage 38.
    assert "Год изобретения" in text
    assert "1837" in text


def test_pptx_extracts_text_and_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "slides.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Лекция: история радио. Изобретатель - Попов."
    tbl = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    tbl.cell(0, 0).text = "Радио"
    tbl.cell(0, 1).text = "Маркони"
    prs.save(str(path))

    text = load_file(str(path))
    assert "история радио" in text
    assert "Попов" in text
    assert "Маркони" in text


def test_rtf_extracts_text(tmp_path):
    path = tmp_path / "doc.rtf"
    # Cyrillic via \'xx escapes (cp1251 = "Привет") + an ASCII word.
    path.write_bytes(rb"{\rtf1\ansi\ansicpg1251 \'cf\'f0\'e8\'e2\'e5\'f2 telegraph}")
    text = load_rtf(str(path))
    assert "telegraph" in text
    assert "Привет" in text


# --- Stage 40: OpenDocument, Excel, CSV -----------------------------------

def test_odt_extracts_paragraphs(tmp_path):
    from odf.opendocument import OpenDocumentText
    from odf.text import H, P

    path = tmp_path / "doc.odt"
    doc = OpenDocumentText()
    h = H(outlinelevel=1)
    h.addText("Заголовок про телеграф")
    doc.text.addElement(h)
    p = P()
    p.addText("Содержимое про радио Попова.")
    doc.text.addElement(p)
    doc.save(str(path))

    text = load_file(str(path))
    assert "Заголовок про телеграф" in text
    assert "радио Попова" in text


def test_ods_extracts_table_cells(tmp_path):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    path = tmp_path / "sheet.ods"
    doc = OpenDocumentSpreadsheet()
    table = Table(name="Лист1")
    row = TableRow()
    for val in ("Год", "1895"):
        cell = TableCell()
        p = P()
        p.addText(val)
        cell.addElement(p)
        row.addElement(cell)
    table.addElement(row)
    doc.spreadsheet.addElement(table)
    doc.save(str(path))

    text = load_file(str(path))
    assert "Год" in text
    assert "1895" in text
    assert "|" in text  # cells joined into a row


def test_xlsx_extracts_rows(tmp_path):
    import openpyxl

    path = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(["Год", "Событие"])
    ws.append([1895, "Радио Попова"])
    wb.save(str(path))

    text = load_file(str(path))
    assert "Радио Попова" in text
    assert "1895" in text
    assert "Год | Событие" in text


def test_csv_extracts_rows(tmp_path):
    path = tmp_path / "table.csv"
    path.write_text("Год;Событие\n1895;Радио Попова\n", encoding="utf-8")

    text = load_file(str(path))
    assert "Радио Попова" in text
    assert "1895" in text
    assert "|" in text


# --- Stage 40: content-based format detection -----------------------------

def test_detects_docx_by_content_when_extension_lies(tmp_path):
    """Файл с НЕВЕРНЫМ расширением читается по содержимому."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Скрытый docx под видом pdf.")
    real = tmp_path / "lecture.pdf"  # wrong extension on purpose
    doc.save(str(real))

    assert detect_format(str(real)) == ".docx"
    text = load_file(str(real))
    assert "Скрытый docx" in text


def test_sniff_pdf_magic_overrides_txt_extension(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_bytes(b"%PDF-1.4\n%fake pdf body\n")
    assert sniff_format(str(path)) == ".pdf"


def test_digital_pdf_extraction_keeps_page_numbers_without_double_read(tmp_path, monkeypatch):
    import pypdf

    class FakePage:
        def __init__(self, text):
            self.text = text
            self.calls = 0

        def extract_text(self):
            self.calls += 1
            return self.text

    pages = [FakePage("Первая страница с правилом."), FakePage("Вторая страница с выводом.")]
    monkeypatch.setattr(pypdf, "PdfReader", lambda _path: type("Reader", (), {"pages": pages})())
    path = tmp_path / "digital.pdf"
    path.write_bytes(b"%PDF-1.4\n%fixture")

    text = load_file(str(path))

    assert "Первая страница" in text and "Вторая страница" in text
    assert document_loader.last_load_pages() == [
        (1, "Первая страница с правилом."),
        (2, "Вторая страница с выводом."),
    ]
    assert [page.calls for page in pages] == [1, 1]


def test_unknown_format_raises(tmp_path):
    path = tmp_path / "blob.xyz"
    path.write_bytes(b"\x01\x02 random bytes, no magic \x03")
    with pytest.raises(UnsupportedFormatError):
        load_file(str(path))


# --- Stage 40: OCR / legacy paths degrade gracefully (no binary present) --

def test_image_ocr_unavailable_friendly_error(tmp_path, monkeypatch):
    from PIL import Image

    path = tmp_path / "scan.png"
    Image.new("RGB", (40, 20), "white").save(str(path))
    monkeypatch.setattr(ocr, "tesseract_available", lambda: False)

    with pytest.raises(ExtractionUnavailableError) as exc:
        load_file(str(path))
    assert "OCR" in str(exc.value)


def test_scanned_pdf_ocr_unavailable_friendly_error(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"%PDF-1.4\n%image-only scan\n")
    # No text layer -> looks like a scan; OCR binary absent.
    monkeypatch.setattr(document_loader, "load_pdf", lambda p: "")
    monkeypatch.setattr(ocr, "tesseract_available", lambda: False)

    with pytest.raises(ExtractionUnavailableError) as exc:
        load_file(str(path))
    msg = str(exc.value)
    assert "скан" in msg and "OCR" in msg


def test_ocr_page_limit_rejects_partial_indexing():
    with pytest.raises(ExtractionUnavailableError) as exc:
        ocr._check_ocr_page_limit(total=120, max_pages=50)

    message = str(exc.value)
    assert "120" in message
    assert "50" in message
    assert "не проиндексирован частично" in message


def test_ocr_page_limit_accepts_complete_document():
    ocr._check_ocr_page_limit(total=50, max_pages=50)


def test_legacy_doc_without_libreoffice_friendly_error(tmp_path, monkeypatch):
    path = tmp_path / "old.doc"
    blob = (
        b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"
        + b"\x00" * 64
        + "WordDocument".encode("utf-16-le")
    )
    path.write_bytes(blob)

    assert detect_format(str(path)) == ".doc"
    monkeypatch.setattr(ocr, "libreoffice_binary", lambda: None)
    with pytest.raises(ExtractionUnavailableError) as exc:
        load_file(str(path))
    assert "LibreOffice" in str(exc.value)


# --- Stage 41: zip-bomb guard + OLE fallback ------------------------------

def test_zip_bomb_guard_rejects_oversized(tmp_path, monkeypatch):
    from docx import Document

    path = tmp_path / "big.docx"
    doc = Document()
    doc.add_paragraph("Содержимое, которое в распакованном виде больше крошечного потолка.")
    doc.save(str(path))

    # Tiny cap so a normal docx trips the guard.
    monkeypatch.setattr(config, "MAX_ZIP_UNCOMPRESSED_BYTES", 10)
    with pytest.raises(ValueError) as exc:
        load_file(str(path))
    assert "распакован" in str(exc.value)


def test_ole_unknown_stream_falls_back_to_extension(tmp_path):
    """OLE-контейнер без опознанного потока: формат решает расширение, поэтому
    настоящий .xls пойдёт в xlrd, а не в LibreOffice (Stage 41)."""
    path = tmp_path / "data.xls"
    path.write_bytes(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1" + b"\x00" * 200)

    assert sniff_format(str(path)) is None
    assert detect_format(str(path)) == ".xls"


# --- .docx: порядок документа и пустые ячейки -----------------------------
#
# Замер 04.09.2026: одна и та же выдержка из свода правил в PDF и в Word.
# Вопрос про коэффициент Kв - по PDF 8 верных из 8, по Word 2 из 4. Ассистент
# писал: «коэффициент можно найти в таблице Ж.3, однако конкретные значения из
# этой таблицы в представленном контексте отсутствуют». Подпись оставалась на
# месте, таблица уезжала в конец файла и попадала в другой фрагмент.


def _docx_with_caption_and_table(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Таблица Ж.3 - Коэффициент взаимодействия")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Число струй"
    table.cell(0, 1).text = "80"
    table.cell(0, 2).text = "100"
    table.cell(1, 0).text = "10"
    table.cell(1, 1).text = "2,4"
    table.cell(1, 2).text = "2,6"
    doc.add_paragraph("После таблицы идёт примечание.")
    path = tmp_path / "caption.docx"
    doc.save(str(path))
    return path


def test_table_follows_its_caption(tmp_path):
    from src.document_loader import load_docx

    text = load_docx(str(_docx_with_caption_and_table(tmp_path)))
    caption = text.index("Таблица Ж.3")
    header = text.index("Число струй")
    after = text.index("После таблицы")
    # Подпись, таблица, примечание - ровно в этом порядке.
    assert caption < header < after


def test_empty_cell_keeps_its_column(tmp_path):
    """«Иванов | (пусто) | 5» не должно схлопываться в «Иванов | 5»."""
    from docx import Document
    from src.document_loader import load_docx

    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Иванов"
    table.cell(0, 2).text = "5"
    table.cell(1, 0).text = "Петров"
    table.cell(1, 1).text = "3"
    table.cell(1, 2).text = "7"
    path = tmp_path / "gaps.docx"
    doc.save(str(path))

    text = load_docx(str(path))
    assert "Иванов |  | 5" in text
    # Обе строки должны иметь одинаковое число разделителей, иначе колонки
    # разъезжаются именно там, где данные разрежены.
    rows = [r for r in text.split("\n") if "|" in r]
    assert len({r.count("|") for r in rows}) == 1


def test_two_tables_do_not_glue_together(tmp_path):
    from docx import Document
    from src.document_loader import load_docx

    doc = Document()
    for value in ("первая", "вторая"):
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = value
        t.cell(0, 1).text = "1"
    path = tmp_path / "two.docx"
    doc.save(str(path))

    text = load_docx(str(path))
    assert "первая | 1\n\nвторая | 1" in text


def test_scan_pdf_keeps_physical_pages_and_resets_for_next_file(tmp_path, monkeypatch):
    scan = tmp_path / "scan.pdf"
    scan.write_bytes(b"%PDF-1.4\n%fixture")
    monkeypatch.setattr(document_loader, "load_pdf", lambda _path: "")
    monkeypatch.setattr(ocr, "ocr_pdf_pages", lambda _path: [(1, "First page."), (3, "Third page.")])
    assert load_file(str(scan)) == "First page.\n\nThird page."
    assert document_loader.last_load_pages() == [(1, "First page."), (3, "Third page.")]
    assert document_loader.last_load_used_ocr()
    text = tmp_path / "next.txt"
    text.write_text("Plain text", encoding="utf-8")
    load_file(str(text))
    assert document_loader.last_load_pages() == []
    assert not document_loader.last_load_used_ocr()


def test_ocr_pdf_does_not_renumber_after_blank_page(monkeypatch):
    import pypdfium2
    import pytesseract
    from PIL import Image

    class Page:
        def render(self, scale):
            return self
        def to_pil(self):
            return Image.new("RGB", (10, 10), "white")
    class PDF:
        def __len__(self):
            return 3
        def __getitem__(self, index):
            return Page()
        def close(self):
            pass
    texts = iter(["First", "   ", "Third"])
    monkeypatch.setattr(ocr, "tesseract_available", lambda: True)
    monkeypatch.setattr(pypdfium2, "PdfDocument", lambda _path: PDF())
    monkeypatch.setattr(
        pytesseract, "image_to_string", lambda image, lang, **kwargs: next(texts)
    )
    assert ocr.ocr_pdf_pages("fixture.pdf") == [(1, "First"), (3, "Third")]


def test_ocr_pdf_rechecks_suspicious_lost_fraction_at_higher_scale(monkeypatch):
    import pypdfium2
    import pytesseract
    from PIL import Image

    scales = []

    class Page:
        def render(self, scale):
            scales.append(scale)
            return self
        def to_pil(self):
            return Image.new("RGB", (10, 10), "white")
    class PDF:
        def __len__(self):
            return 1
        def __getitem__(self, index):
            return Page()
        def close(self):
            pass

    texts = iter(["Drink 34-1 quart per hour", "Drink 3/4-1 quart per hour"])
    monkeypatch.setattr(ocr, "tesseract_available", lambda: True)
    monkeypatch.setattr(pypdfium2, "PdfDocument", lambda _path: PDF())
    monkeypatch.setattr(
        pytesseract, "image_to_string",
        lambda image, lang, **kwargs: next(texts),
    )

    assert ocr.ocr_pdf_pages("fixture.pdf") == [(1, "Drink 3/4-1 quart per hour")]
    assert scales == [ocr.OCR_PDF_SCALE, ocr.OCR_PDF_RECHECK_SCALE]


def test_ocr_pdf_can_use_vision_only_after_both_local_passes_fail(monkeypatch):
    import pypdfium2
    import pytesseract
    from PIL import Image

    class Page:
        def render(self, scale):
            return self
        def to_pil(self):
            return Image.new("RGB", (10, 10), "white")
    class PDF:
        def __len__(self):
            return 1
        def __getitem__(self, index):
            return Page()
        def close(self):
            pass

    texts = iter(["This translates to 34-1 quart", "This translates to 34-1 quart"])
    vision = "This translates to 3/4-1 quart. " + ("Verified page text. " * 8)
    monkeypatch.setattr(ocr, "tesseract_available", lambda: True)
    monkeypatch.setattr(ocr, "OCR_VISION_RECHECK_ENABLED", True)
    monkeypatch.setattr(ocr, "_vision_ocr_page", lambda page: vision)
    monkeypatch.setattr(pypdfium2, "PdfDocument", lambda _path: PDF())
    monkeypatch.setattr(
        pytesseract, "image_to_string", lambda image, lang, **kwargs: next(texts)
    )

    pages = ocr.ocr_pdf_pages("fixture.pdf")
    assert pages == [(1, vision)]
    assert "[OCR:" not in pages[0][1]
